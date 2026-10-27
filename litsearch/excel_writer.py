# Litwell, Copyright (c) 2026 J. Kent Morgan Consulting, LLC. Created by J. Kent Morgan, Ph.D.
# Licensed under the Apache License, Version 2.0; see the LICENSE file.
"""The output workbook: Results, Summary, Search Log and Glossary sheets.

The workbook is rebuilt on every run from three things: the records already in
the file, the new records, and the log. That keeps Results sorted best-first,
the Summary current, and lets a record found again in a later run be merged
(extra sources, missing fields) instead of duplicated.

Files made by older versions (a Results sheet starting with Title, one Source
text column) are read and upgraded on the next run.
"""

from __future__ import annotations

import datetime as _dt
import os
import shutil
import zipfile
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import CREDIT, LICENSE_LINE, __version__
from .models import Article, SearchParams, make_dedupe_key
from .records import sort_best_first

RESULTS_SHEET = "Results"
SUMMARY_SHEET = "Summary"
LOG_SHEET = "Search Log"
GLOSSARY_SHEET = "Glossary"

# source key -> (column header, name older versions wrote in the Source column)
SOURCE_COLUMNS: Dict[str, Tuple[str, Tuple[str, ...]]] = {
    "pubmed": ("PubMed", ("PubMed",)),
    "crossref": ("Crossref", ("Crossref",)),
    "openalex": ("OpenAlex", ("OpenAlex",)),
    "europepmc": ("Europe PMC", ("Europe PMC",)),
    "arxiv": ("arXiv", ("arXiv",)),
    "preprints": ("Preprint servers", ("Preprint",)),
    "patents": ("USPTO", ("USPTO Patents", "USPTO")),
    "epo": ("EPO", ("EPO / Worldwide Patents", "EPO")),
}

RESULTS_HEADERS = (
    ["Relevance Score", "AI Rating", "AI Reason", "New Since Last Run",
     "Title", "Authors", "Journal", "Date", "Abstract", "DOI", "Link",
     "Open-access link", "Citation count", "Type"]
    + [h for h, _ in SOURCE_COLUMNS.values()]
    + ["Search Date", "Search Terms"]
)
LOG_HEADERS = [
    "Run At", "Source", "Status", "Keyword", "Match Mode", "Must Also Contain",
    "Exclude Terms", "Author", "Title", "Journal", "Date Start", "Date End",
    "Matches Reported", "Records Fetched", "Rows Removed", "New Rows Added",
    "Duplicates Skipped", "Details",
]

_HEADER_FILL = PatternFill("solid", fgColor="001F5B")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single")
_SECTION_FONT = Font(bold=True, size=12, color="001F5B")
_WIDTHS = {
    "Relevance Score": 11, "AI Rating": 13, "AI Reason": 40, "New Since Last Run": 11,
    "Title": 60, "Authors": 34, "Journal": 28, "Date": 12, "Abstract": 90,
    "DOI": 26, "Link": 38, "Open-access link": 38, "Citation count": 10, "Type": 10,
    "Search Date": 19, "Search Terms": 44,
}
_YES_NO_WIDTH = 9
_ABSTRACT_ROW_HEIGHT = 78


class DamagedWorkbook(RuntimeError):
    """Raised once when an unreadable workbook is moved aside. Retrying the
    same search will then succeed against the fresh file."""


# ---------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------
def _header_map(ws) -> Dict[str, int]:
    return {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}


def _open(path: Path) -> Optional["Workbook"]:
    if path.exists() and not path.is_file():
        raise RuntimeError(f"'{path.name}' is a folder, not a workbook file.")
    if not (path.exists() and path.stat().st_size > 0):
        return None
    try:
        return load_workbook(path)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        stamp = _dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        damaged = path.with_name(f"{path.stem} (damaged {stamp}){path.suffix}")
        try:
            path.rename(damaged)
        except OSError:
            damaged = None
        raise DamagedWorkbook(
            f"'{path.name}' could not be opened ({exc.__class__.__name__}). "
            + (f"It was moved aside as '{damaged.name}' and a new file started."
               if damaged else "A new file was started.")
        )


def _cell(row, hm: Dict[str, int], name: str):
    i = hm.get(name)
    return row[i] if i is not None and i < len(row) else None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def read_records(ws) -> List[Article]:
    """Rows of a Results sheet (this version's or an older layout) as Articles."""
    hm = _header_map(ws)
    out: List[Article] = []
    by_header = {v[0]: k for k, v in SOURCE_COLUMNS.items()}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        a = Article(title=_text(_cell(row, hm, "Title")))
        a.authors = [x.strip() for x in _text(_cell(row, hm, "Authors")).split(";") if x.strip()]
        a.journal = _text(_cell(row, hm, "Journal"))
        a.date = _text(_cell(row, hm, "Date"))
        a.abstract = _text(_cell(row, hm, "Abstract"))
        a.doi = _text(_cell(row, hm, "DOI"))
        a.link = _text(_cell(row, hm, "Link"))
        a.oa_url = _text(_cell(row, hm, "Open-access link"))
        c = _cell(row, hm, "Citation count")
        a.citations = int(c) if isinstance(c, (int, float)) else None
        a.kind = _text(_cell(row, hm, "Type"))
        s = _cell(row, hm, "Relevance Score")
        a.score = int(s) if isinstance(s, (int, float)) else None
        a.ai_rating = _text(_cell(row, hm, "AI Rating"))
        a.ai_reason = _text(_cell(row, hm, "AI Reason"))
        a.search_date = _text(_cell(row, hm, "Search Date"))
        a.search_terms = _text(_cell(row, hm, "Search Terms"))
        for header, key in by_header.items():
            if _text(_cell(row, hm, header)).lower() == "yes":
                a.sources.append(key)
        old = _text(_cell(row, hm, "Source"))          # layout of versions before 1.4
        if old and not a.sources:
            a.source = old
            for key, (_h, names) in SOURCE_COLUMNS.items():
                if any(n.lower() in old.lower() for n in names):
                    a.sources.append(key)
        if not a.kind:
            a.kind = _kind_from_sources(a.sources)
        out.append(a)
    return out


def _kind_from_sources(sources: Sequence[str]) -> str:
    if sources and all(s in ("patents", "epo") for s in sources):
        return "Patent"
    if sources and all(s in ("arxiv", "preprints") for s in sources):
        return "Preprint"
    return "Paper" if sources else ""


def read_log(wb) -> List[List]:
    """Existing Search Log rows, mapped onto this version's columns by name."""
    if LOG_SHEET not in wb.sheetnames:
        return []
    ws = wb[LOG_SHEET]
    hm = _header_map(ws)
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        mapped = [_cell(row, hm, h) for h in LOG_HEADERS]
        if not _text(mapped[2]):                       # an older log had no Status column
            mapped[2] = "OK"
        rows.append(mapped)
    return rows


def load_existing(path_str: str) -> Tuple[List[Article], List[List], bool]:
    """(records, log rows, file_existed) for the output file."""
    path = Path(path_str).expanduser()
    wb = _open(path)
    if wb is None:
        return [], [], False
    arts = read_records(wb[RESULTS_SHEET]) if RESULTS_SHEET in wb.sheetnames else []
    return arts, read_log(wb), True


# ---------------------------------------------------------------------------
# writing
# ---------------------------------------------------------------------------
GLOSSARY: List[Tuple[str, str]] = [
    ("Relevance Score", "0 to 100, best first. Up to 50 points when your topic appears as a phrase in the title (otherwise up to 30 for the share of its words in the title); up to 25 when it is in the abstract (otherwise up to 15 for the share of words); 10 per 'Must also contain' term found, up to 20; 5 for a journal or field match; minus 20 per exclusion term found only in the abstract, up to 40. It is a rule-based guide, not a judgement of quality."),
    ("AI Rating", "Relevant, Maybe or Not relevant, from an optional local Ollama model judging the title and abstract against the profile description. Blank when AI triage was off or unavailable."),
    ("AI Reason", "One line from the local model explaining its rating."),
    ("New Since Last Run", "Yes when the record was not in this file before the latest run; No when it was. Blank for the first run into a new file."),
    ("Title / Authors / Journal / Date", "As published. For patents, Authors are the inventors and Journal is the applicant."),
    ("Abstract", "The abstract when a source supplied one (PubMed, Europe PMC, OpenAlex, arXiv, preprint servers, EPO, and Crossref when present). Blank otherwise."),
    ("DOI", "Digital Object Identifier. For patents this column holds the patent or publication number."),
    ("Link", "A page for the record (publisher, PubMed, Google Patents ...)."),
    ("Open-access link", "A free full-text link, when a source reports one."),
    ("Citation count", "How many works cite this one, from OpenAlex, Crossref or Europe PMC (the highest figure found). Counts differ between databases."),
    ("Type", "Paper, Preprint (not yet peer reviewed) or Patent."),
    ("PubMed ... EPO columns", "Yes when that source returned the record. A paper found by several sources is one row with several Yes."),
    ("Search Date / Search Terms", "When the record was first found and the fields used; a later search that finds it again adds its terms."),
    ("Exact phrase", "The topic is searched as one phrase, so 'B2 protein' does not match a paper that only has 'B2' and 'protein' far apart."),
    ("All words / Any word", "Every word must appear (any order) / at least one word must appear."),
    ("Must also contain", "At least one of these terms must appear in the title or abstract, or the record is removed."),
    ("Exclude terms", "A term in the title or journal removes the record; a term only in the abstract lowers its score."),
    ("Preprint", "A paper shared before peer review (bioRxiv, arXiv ...)."),
    ("Summary sheet", "Counts of what the file holds, by source, type, year, AI rating and score band."),
    ("Search Log sheet", "One row per source per run (including errors and skipped sources) and one row per filter rule with the number of rows it removed."),
]


def _style_header(ws, headers: Sequence[str], widths: Dict[str, int]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 20)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_results(ws, arts: List[Article]) -> None:
    widths = dict(_WIDTHS)
    for h, _ in SOURCE_COLUMNS.values():
        widths[h] = _YES_NO_WIDTH
    _style_header(ws, RESULTS_HEADERS, widths)
    col = {h: i for i, h in enumerate(RESULTS_HEADERS, start=1)}
    center = Alignment(horizontal="center", vertical="top")
    top = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)
    for r, a in enumerate(arts, start=2):
        ws.cell(r, col["Relevance Score"], a.score).alignment = center
        ws.cell(r, col["AI Rating"], a.ai_rating or None).alignment = center
        ws.cell(r, col["AI Reason"], a.ai_reason or None).alignment = wrap
        ws.cell(r, col["New Since Last Run"],
                {True: "Yes", False: "No"}.get(a.is_new)).alignment = center
        ws.cell(r, col["Title"], a.title).alignment = wrap
        ws.cell(r, col["Authors"], a.authors_str).alignment = top
        ws.cell(r, col["Journal"], a.journal).alignment = top
        ws.cell(r, col["Date"], a.date).alignment = top
        ws.cell(r, col["Abstract"], a.abstract or None).alignment = wrap
        ws.cell(r, col["DOI"], a.doi).alignment = top
        link = ws.cell(r, col["Link"], a.link or a.doi)
        link.alignment = top
        if a.link:
            link.hyperlink = a.link
            link.font = _LINK_FONT
        oa = ws.cell(r, col["Open-access link"], a.oa_url or None)
        oa.alignment = top
        if a.oa_url:
            oa.hyperlink = a.oa_url
            oa.font = _LINK_FONT
        ws.cell(r, col["Citation count"], a.citations).alignment = center
        ws.cell(r, col["Type"], a.kind or None).alignment = center
        for key, (header, _n) in SOURCE_COLUMNS.items():
            ws.cell(r, col[header], "Yes" if key in a.sources else "No").alignment = center
        ws.cell(r, col["Search Date"], a.search_date).alignment = top
        ws.cell(r, col["Search Terms"], a.search_terms).alignment = wrap
        if a.abstract:
            ws.row_dimensions[r].height = _ABSTRACT_ROW_HEIGHT
    if arts:
        letter = get_column_letter(col["Relevance Score"])
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{len(arts) + 1}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8D7DA",
                           mid_type="num", mid_value=50, mid_color="FFF3CD",
                           end_type="num", end_value=100, end_color="D1E7DD"),
        )


def _year(date: str) -> str:
    y = (date or "")[:4]
    return y if y.isdigit() else "Unknown"


def _band(score: Optional[int]) -> str:
    if score is None:
        return "Not scored"
    for lo, label in ((80, "80 to 100"), (60, "60 to 79"), (40, "40 to 59"), (20, "20 to 39")):
        if score >= lo:
            return label
    return "0 to 19"


def _write_summary(ws, arts: List[Article], run_info: Dict) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.cell(1, 1, "Litwell summary").font = Font(bold=True, size=14, color="001F5B")
    r = 3

    def section(title: str, rows: Sequence[Tuple[str, object]]) -> None:
        nonlocal r
        ws.cell(r, 1, title).font = _SECTION_FONT
        r += 1
        for label, value in rows:
            ws.cell(r, 1, label)
            ws.cell(r, 2, value).alignment = Alignment(horizontal="right")
            r += 1
        r += 1

    section("This file", [
        ("Records in Results", len(arts)),
        ("Latest run", run_info.get("run_at", "")),
        ("Rows added in latest run", run_info.get("added", "")),
        ("Rows removed by filters in latest run", run_info.get("removed", "")),
    ])
    section("By source (a record can have several)",
            [(h, sum(1 for a in arts if k in a.sources)) for k, (h, _n) in SOURCE_COLUMNS.items()])
    kinds = Counter(a.kind or "Unknown" for a in arts)
    section("By type", [(k, kinds.get(k, 0)) for k in ("Paper", "Preprint", "Patent", "Unknown") if kinds.get(k, 0)])
    years = Counter(_year(a.date) for a in arts)
    section("By year", sorted(years.items(), key=lambda kv: (kv[0] == "Unknown", kv[0]), reverse=False))
    ratings = Counter(a.ai_rating or "Not rated" for a in arts)
    section("By AI rating", [(k, ratings.get(k, 0)) for k in ("Relevant", "Maybe", "Not relevant", "Not rated") if ratings.get(k, 0)])
    bands = Counter(_band(a.score) for a in arts)
    section("By relevance score", [(k, bands.get(k, 0)) for k in ("80 to 100", "60 to 79", "40 to 59", "20 to 39", "0 to 19", "Not scored") if bands.get(k, 0)])
    ws.cell(r, 1, CREDIT).font = Font(italic=True, color="666666")


def _write_log(ws, rows: List[List]) -> None:
    widths = {"Run At": 19, "Source": 24, "Status": 34, "Details": 60, "Keyword": 28,
              "Must Also Contain": 28, "Exclude Terms": 28}
    _style_header(ws, LOG_HEADERS, widths)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v).alignment = Alignment(vertical="top", wrap_text=c in (3, 18))


def _write_glossary(ws) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    for c, name in enumerate(("Term", "Meaning"), start=1):
        cell = ws.cell(1, c, name)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    for r, (term, meaning) in enumerate(GLOSSARY, start=2):
        ws.cell(r, 1, term).font = Font(bold=True)
        ws.cell(r, 2, meaning).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
    foot = len(GLOSSARY) + 3
    ws.cell(foot, 1, f"Made with Litwell {__version__}").font = Font(italic=True, color="666666")
    ws.cell(foot + 1, 1, CREDIT).font = Font(italic=True, color="666666")
    ws.cell(foot + 2, 1, LICENSE_LINE).font = Font(italic=True, color="666666")
    ws.freeze_panes = "A2"


def _print_setup(ws) -> None:
    """Landscape, one page wide, header row repeated: prints and exports to PDF sensibly."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build_workbook(arts: List[Article], log_rows: List[List], run_info: Dict) -> Workbook:
    wb = Workbook()
    results = wb.active
    results.title = RESULTS_SHEET
    _write_results(results, arts)
    _write_summary(wb.create_sheet(SUMMARY_SHEET), arts, run_info)
    _write_log(wb.create_sheet(LOG_SHEET), log_rows)
    _write_glossary(wb.create_sheet(GLOSSARY_SHEET))
    for ws in wb.worksheets:
        _print_setup(ws)
    return wb


def save_run(
    path_str: str,
    existing: List[Article],
    existing_log: List[List],
    new_articles: List[Article],
    new_log_rows,
    run_info: Dict,
) -> Tuple[int, int]:
    """Merge ``new_articles`` into ``existing`` and rewrite the workbook.

    Returns (added, already_known). ``existing`` are records from earlier runs
    (empty for a new file); a new record found again is merged into its row.
    """
    from .records import Merger

    path = Path(path_str).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    had_file = bool(existing) or bool(existing_log)

    m = Merger()
    for a in existing:
        a.is_new = False if had_file else None
        m.add(a)
    added = known = 0
    for a in new_articles:
        if m.add(a):
            a.is_new = True if had_file else None
            added += 1
        else:
            a.is_new = False             # already in the file: its row stays "No"
            known += 1
    merged = sort_best_first(m.items)

    run_info = dict(run_info, added=added)
    rows = new_log_rows() if callable(new_log_rows) else new_log_rows
    wb = build_workbook(merged, existing_log + rows, run_info)
    _save(wb, path)
    return added, known


# ---------------------------------------------------------------------------
# tools used by the app
# ---------------------------------------------------------------------------
def _scan_ws_duplicates(ws) -> List[Tuple[int, str, int]]:
    """(row, title, first_seen_row) for every duplicate row, keeping the first."""
    hm = _header_map(ws)
    ti, di = hm.get("Title", 0), hm.get("DOI", 5)
    seen: Dict[str, int] = {}
    dupes: List[Tuple[int, str, int]] = []
    for r in range(2, ws.max_row + 1):
        title = ws.cell(row=r, column=ti + 1).value
        doi = ws.cell(row=r, column=di + 1).value
        if title is None and doi is None:
            continue
        key = make_dedupe_key(str(title or ""), str(doi or ""))
        if not key:
            continue
        if key in seen:
            dupes.append((r, str(title or "(untitled)"), seen[key]))
        else:
            seen[key] = r
    return dupes


def review_duplicates(path_str: str) -> List[Tuple[str, int]]:
    path = Path(path_str).expanduser()
    if not path.exists():
        raise RuntimeError("The Excel file does not exist yet. Run a search first.")
    wb = load_workbook(path)
    if RESULTS_SHEET not in wb.sheetnames:
        return []
    return [(t, r) for (r, t, _first) in _scan_ws_duplicates(wb[RESULTS_SHEET])]


def remove_duplicates(path_str: str, *, make_backup: bool = True) -> Dict[str, object]:
    """Delete duplicate rows (keeping the first of each) after saving a backup copy."""
    path = Path(path_str).expanduser()
    if not path.exists():
        raise RuntimeError("The Excel file does not exist yet. Run a search first.")
    wb = load_workbook(path)
    if RESULTS_SHEET not in wb.sheetnames:
        return {"removed": 0, "backup": None, "titles": []}
    ws = wb[RESULTS_SHEET]
    dupes = _scan_ws_duplicates(ws)
    if not dupes:
        return {"removed": 0, "backup": None, "titles": []}
    backup: Optional[Path] = None
    if make_backup:
        stamp = _dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        backup = path.with_name(f"{path.stem} (backup {stamp}){path.suffix}")
        shutil.copy2(path, backup)
    for row_num in sorted((d[0] for d in dupes), reverse=True):
        ws.delete_rows(row_num, 1)
    _save(wb, path)
    return {"removed": len(dupes), "backup": str(backup) if backup else None,
            "titles": [d[1] for d in dupes]}


def count_entries(path_str: str) -> int:
    """Number of result rows already in the workbook (0 if missing/unreadable)."""
    path = Path(path_str).expanduser()
    if not path.exists() or path.stat().st_size == 0:
        return 0
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:
        return 0
    if RESULTS_SHEET not in wb.sheetnames:
        return 0
    return sum(1 for row in wb[RESULTS_SHEET].iter_rows(min_row=2, values_only=True)
               if row and any(row))


def _save(wb: Workbook, path: Path) -> None:
    """Save atomically: write a sibling temp file, then swap it in. A failure
    (file open in Excel, disk full, permissions) leaves the original intact."""
    tmp = path.with_name(f".~{path.stem}.{os.getpid()}.tmp")
    try:
        wb.save(tmp)
    except PermissionError as exc:
        _quiet_unlink(tmp)
        raise RuntimeError(
            f"Could not write to the folder for '{path.name}'. "
            f"Check the location is not read-only."
        ) from exc
    except OSError as exc:
        _quiet_unlink(tmp)
        raise RuntimeError(f"Could not save '{path.name}': {exc}") from exc
    try:
        os.replace(tmp, path)
    except PermissionError as exc:
        _quiet_unlink(tmp)
        raise RuntimeError(
            f"Could not update '{path.name}': it is open in Excel "
            f"(or another program). Close it and search again."
        ) from exc
    except OSError as exc:
        _quiet_unlink(tmp)
        raise RuntimeError(f"Could not update '{path.name}': {exc}") from exc


def _quiet_unlink(p: Path) -> None:
    try:
        p.unlink()
    except OSError:
        pass
