# Litwell, Copyright (c) 2026 J. Kent Morgan Consulting, LLC. Created by J. Kent Morgan, Ph.D.
# Licensed under the Apache License, Version 2.0; see the LICENSE file.
"""Offline tests for searching several sources at once (no network).

Run from the project folder:   python -m unittest discover -s tests -v
"""

import os
import sys
import tempfile
import threading
import time
import types
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from litsearch import output, search                      # noqa: E402
from litsearch.config import selected_sources             # noqa: E402
from litsearch.http_client import NetworkError            # noqa: E402
from litsearch.models import Article, SearchParams        # noqa: E402


def art(title, doi="", source="x"):
    a = Article(source=source)
    a.title, a.doi = title, doi
    a.authors, a.date = ["A. Author"], "2025-01-01"
    return a


def fake(records=None, error=None, delay=0.0, log=None):
    """A stand-in source module with the same search() signature."""
    def _search(params, cfg):
        if log is not None:
            log.append((params.source, threading.current_thread().name))
        time.sleep(delay)
        if error:
            raise error
        return list(records or []), len(records or [])
    return types.SimpleNamespace(search=_search)


class MultiSearchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.cfg = {"excel_path": os.path.join(self.tmp, "out.csv")}
        self.params = SearchParams(keyword="crispr", max_results=10)
        self._orig = dict(search.SOURCES)

    def tearDown(self):
        search.SOURCES.clear()
        search.SOURCES.update(self._orig)

    def use(self, **sources):
        search.SOURCES.clear()
        for k, mod in sources.items():
            search.SOURCES[k] = (k.upper(), mod)

    def test_two_sources_saved_once_and_cross_source_duplicate_skipped(self):
        self.use(a=fake([art("Paper One", "10.1234/one"), art("Paper Two", "10.1234/two")]),
                 b=fake([art("Paper Two again", "10.1234/TWO"), art("Paper Three", "10.1234/three")]))
        m = search.run_multi_search(self.params, ["a", "b"], self.cfg)
        self.assertTrue(m.ok)
        self.assertEqual([r.key for r in m.results], ["a", "b"])
        self.assertEqual(m.results[0].outcome.added, 2)
        self.assertEqual(m.results[1].outcome.added, 2)             # Paper Two (shared) + Paper Three
        self.assertEqual(m.merged_copies, 1)                        # same DOI, other case, merged
        self.assertEqual(m.added, 3)
        self.assertEqual(output.count_entries(self.cfg["excel_path"]), 3)

    def test_one_failing_source_does_not_stop_the_others(self):
        self.use(a=fake([art("Paper One", "10.1234/one")]),
                 b=fake(error=ValueError("USPTO patent search needs a free API key.")),
                 c=fake([art("Paper Three", "10.1234/three")]))
        m = search.run_multi_search(self.params, ["a", "b", "c"], self.cfg)
        self.assertTrue(m.ok)
        self.assertEqual([r.key for r in m.succeeded], ["a", "c"])
        self.assertEqual([r.key for r in m.failed], ["b"])
        self.assertIn("free API key", m.failed[0].outcome.error)
        self.assertEqual(output.count_entries(self.cfg["excel_path"]), 2)

    def test_network_error_and_unexpected_error_are_captured(self):
        self.use(a=fake(error=NetworkError("rate limit reached")),
                 b=fake(error=RuntimeError("boom")))
        m = search.run_multi_search(self.params, ["a", "b"], self.cfg)
        self.assertFalse(m.ok)                       # nothing succeeded
        self.assertEqual(len(m.failed), 2)
        self.assertIn("rate limit", m.failed[0].outcome.error)
        self.assertIn("boom", m.failed[1].outcome.error)

    def test_file_order_follows_the_requested_order_not_finish_order(self):
        self.use(slow=fake([art("From slow", "10.1/slow")], delay=0.4),
                 fast=fake([art("From fast", "10.1/fast")]))
        search.run_multi_search(self.params, ["slow", "fast"], self.cfg)
        text = Path(self.cfg["excel_path"]).read_text(encoding="utf-8")
        self.assertLess(text.index("From slow"), text.index("From fast"))

    def test_sources_run_in_parallel(self):
        log = []
        self.use(**{k: fake([art(k, f"10.1/{k}")], delay=0.5, log=log) for k in "abcd"})
        t0 = time.time()
        search.run_multi_search(self.params, list("abcd"), self.cfg)
        self.assertLess(time.time() - t0, 1.4)       # 4 x 0.5 s if sequential = 2 s
        self.assertGreater(len({thread for _k, thread in log}), 1)

    def test_validation_errors_do_nothing(self):
        self.use(a=fake([art("x", "10.1/x")]))
        for sources, params, cfg, needle in [
            (["nope"], self.params, self.cfg, "Unknown source"),
            ([], self.params, self.cfg, "at least one source"),
            (["a"], SearchParams(), self.cfg, "at least one search field"),
            (["a"], self.params, {"excel_path": ""}, "output file"),
        ]:
            m = search.run_multi_search(params, sources, cfg)
            self.assertIn(needle, m.error)
            self.assertFalse(m.ok)
            self.assertEqual(m.results, [])
        self.assertFalse(os.path.exists(self.cfg["excel_path"]))

    def test_duplicate_source_keys_are_searched_once(self):
        log = []
        self.use(a=fake([art("One", "10.1/one")], log=log))
        m = search.run_multi_search(self.params, ["a", "A", "a"], self.cfg)
        self.assertEqual(len(m.results), 1)
        self.assertEqual(len(log), 1)

    def test_unwritable_file_gives_one_recovery_file_and_keeps_every_record(self):
        self.use(a=fake([art("One", "10.1/one")]), b=fake([art("Two", "10.1/two")]),
                 c=fake([art("Three", "10.1/three")]))
        # A directory where the output file should be makes every write fail.
        os.mkdir(self.cfg["excel_path"])
        m = search.run_multi_search(self.params, ["a", "b", "c"], self.cfg)
        self.assertFalse(m.ok)
        recovered = [f for f in os.listdir(self.tmp) if "recovered" in f]
        self.assertGreaterEqual(len(recovered), 1)
        text = "".join(Path(self.tmp, f).read_text(encoding="utf-8") for f in recovered)
        for title in ("One", "Two", "Three"):
            self.assertIn(title, text)

    def test_single_run_search_still_works(self):
        self.use(a=fake([art("One", "10.1/one")]))
        o = search.run_search(SearchParams(keyword="x", source="a"), self.cfg)
        self.assertTrue(o.ok)
        self.assertEqual(o.added, 1)

    def test_selected_sources_migrates_the_old_single_setting(self):
        valid = ["pubmed", "crossref", "arxiv"]
        self.assertEqual(selected_sources({"default_source": "arxiv"}, valid), ["arxiv"])
        self.assertEqual(selected_sources({"default_sources": ["crossref", "pubmed"],
                                           "default_source": "arxiv"}, valid), ["crossref", "pubmed"])
        self.assertEqual(selected_sources({"default_sources": ["gone"], "default_source": "gone"}, valid),
                         ["pubmed"])


class WorkbookTests(unittest.TestCase):
    """The Excel output: columns, sheets, merging across sources and runs."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = os.path.join(self.tmp, "r.xlsx")
        self.cfg = {"excel_path": self.path}
        self._orig = dict(search.SOURCES)

    def tearDown(self):
        search.SOURCES.clear()
        search.SOURCES.update(self._orig)

    def use(self, **sources):
        search.SOURCES.clear()
        for k, mod in sources.items():
            search.SOURCES[k] = (k.upper(), mod)

    def sheet(self, name="Results"):
        from openpyxl import load_workbook
        ws = load_workbook(self.path)[name]
        head = [c.value for c in ws[1]]
        return head, [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    def test_four_sheets_and_the_requested_results_columns(self):
        from litsearch import excel_writer as ew
        self.use(pubmed=fake([art("Alpha beta gamma study", "10.1/a")]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["Results", "Summary", "Search Log", "Glossary"])
        head, rows = self.sheet()
        for col in ("Relevance Score", "AI Rating", "AI Reason", "Title", "Authors", "Journal", "Date",
                    "Abstract", "DOI", "Link", "Open-access link", "Citation count", "Type",
                    "PubMed", "Crossref", "OpenAlex", "Europe PMC", "arXiv", "Preprint servers",
                    "USPTO", "EPO", "Search Date", "Search Terms"):
            self.assertIn(col, head)
        self.assertEqual(rows[0]["PubMed"], "Yes")
        self.assertEqual(rows[0]["Crossref"], "No")
        self.assertEqual(wb["Results"].freeze_panes, "A2")
        self.assertTrue(wb["Results"].auto_filter.ref)

    def test_a_paper_found_by_two_sources_is_one_row_with_both_flags_and_gaps_filled(self):
        a = art("Shared paper title here", "10.1/p", "PubMed")
        a.abstract = ""
        a.authors = []
        b = art("Shared paper title here", "10.1/P", "Crossref")
        b.abstract, b.citations, b.oa_url = "An abstract about alpha.", 7, "https://example.org/oa"
        b.authors = ["Z. Real"]
        self.use(pubmed=fake([a]), crossref=fake([b]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed", "crossref"], self.cfg)
        _h, rows = self.sheet()
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["PubMed"], r["Crossref"]), ("Yes", "Yes"))
        self.assertEqual(r["Authors"], "Z. Real")                  # filled from the other copy
        self.assertEqual(r["Abstract"], "An abstract about alpha.")
        self.assertEqual((r["Citation count"], r["Open-access link"]), (7, "https://example.org/oa"))

    def test_missing_doi_is_filled_from_a_copy_found_by_title(self):
        a = art("A reasonably long shared title", "", "PubMed")
        b = art("A reasonably long shared title", "10.1/found", "Crossref")
        self.use(pubmed=fake([a]), crossref=fake([b]))
        search.run_multi_search(SearchParams(keyword="shared title"), ["pubmed", "crossref"], self.cfg)
        _h, rows = self.sheet()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["DOI"], "10.1/found")

    def test_new_since_last_run_flag(self):
        self.use(pubmed=fake([art("First paper about alpha", "10.1/one")]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        _h, rows = self.sheet()
        self.assertIsNone(rows[0]["New Since Last Run"])            # first run into a new file
        self.use(pubmed=fake([art("First paper about alpha", "10.1/one"),
                              art("Second paper about alpha", "10.1/two")]))
        m = search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        self.assertEqual((m.new_rows, m.known_rows, m.total_rows), (1, 1, 2))
        _h, rows = self.sheet()
        self.assertEqual({r["Title"]: r["New Since Last Run"] for r in rows},
                         {"First paper about alpha": "No", "Second paper about alpha": "Yes"})

    def test_results_are_sorted_best_first(self):
        self.use(pubmed=fake([art("Unrelated thing", "10.1/u"), art("Alpha beta protein study", "10.1/a")]))
        search.run_multi_search(SearchParams(keyword="alpha beta"), ["pubmed"], self.cfg)
        _h, rows = self.sheet()
        self.assertEqual(rows[0]["Title"], "Alpha beta protein study")
        self.assertGreater(rows[0]["Relevance Score"], rows[1]["Relevance Score"])

    def test_file_from_an_older_version_is_upgraded_and_keeps_its_data(self):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "Results"
        ws.append(["Title", "Authors", "Journal", "Date", "Link", "DOI", "Search Date", "Search Terms", "Source"])
        ws.append(["Old paper about alpha", "B. Old", "J Old", "2020", "", "10.1/old", "2025-01-01 10:00:00", "Keyword: x", "PubMed; Crossref"])
        log = wb.create_sheet("Search Log")
        log.append(["Run At", "Source", "Keyword", "Author", "Title", "Journal", "Date Start", "Date End",
                   "Matches Reported", "Records Fetched", "New Rows Added", "Duplicates Skipped"])
        log.append(["2025-01-01 10:00:00", "pubmed", "x", None, None, None, None, None, 5, 5, 5, 0])
        wb.save(self.path)
        self.use(arxiv=fake([art("New paper about alpha", "10.1/new")]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["arxiv"], self.cfg)
        _h, rows = self.sheet()
        by = {r["Title"]: r for r in rows}
        old = by["Old paper about alpha"]
        self.assertEqual((old["PubMed"], old["Crossref"], old["arXiv"]), ("Yes", "Yes", "No"))
        self.assertEqual(old["Search Date"], "2025-01-01 10:00:00")
        self.assertEqual(old["New Since Last Run"], "No")
        self.assertEqual(by["New paper about alpha"]["New Since Last Run"], "Yes")
        _lh, logrows = self.sheet("Search Log")
        self.assertEqual(logrows[0]["Keyword"], "x")                # the old log row survives
        self.assertGreaterEqual(len(logrows), 2)

    def test_search_log_lists_every_source_errors_skips_and_filter_counts(self):
        self.use(pubmed=fake([art("Good alpha paper", "10.1/g"), art("Vitamin alpha paper", "10.1/v")]),
                 crossref=fake(error=NetworkError("rate limit reached")),
                 patents=fake([art("never called", "10.1/n")]))
        cfg = dict(self.cfg)            # no USPTO key -> patents is skipped
        p = SearchParams(keyword="alpha", exclude_terms="vitamin")
        m = search.run_multi_search(p, ["pubmed", "crossref", "patents"], cfg)
        self.assertTrue(m.ok)
        self.assertEqual(m.rules, [("Exclude: vitamin", 1)])
        _h, rows = self.sheet("Search Log")
        status = {r["Source"]: r["Status"] for r in rows}
        self.assertEqual(status["PUBMED"], "OK")
        self.assertTrue(status["CROSSREF"].startswith("Error"))
        self.assertEqual(status["PATENTS"], "Skipped: no key")
        rule = [r for r in rows if r["Status"] == "Exclude: vitamin"][0]
        self.assertEqual(rule["Rows Removed"], 1)

    def test_summary_and_glossary_have_the_credit(self):
        self.use(pubmed=fake([art("Alpha paper", "10.1/a")]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        text = " ".join(str(c.value) for ws in (wb["Summary"], wb["Glossary"]) for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("Copyright (c) 2026 J. Kent Morgan Consulting, LLC", text)
        self.assertIn("Apache License", text)
        self.assertIn("Relevance Score", text)

    def test_duplicate_tool_still_works_on_the_new_layout(self):
        from litsearch import excel_writer as ew
        self.use(pubmed=fake([art("Alpha paper", "10.1/a")]))
        search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        self.assertEqual(ew.review_duplicates(self.path), [])
        self.assertEqual(ew.count_entries(self.path), 1)

    def test_locked_file_gives_a_recovery_csv(self):
        self.use(pubmed=fake([art("Alpha paper", "10.1/a")]))
        os.mkdir(self.path)                                   # cannot be a workbook
        m = search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], self.cfg)
        self.assertFalse(m.ok)
        self.assertTrue([f for f in os.listdir(self.tmp) if "recovered" in f])


class MatchModeTests(unittest.TestCase):
    """Each source builds its own native phrase / all / any syntax."""

    def p(self, mode):
        return SearchParams(keyword="B2 Protein", match_mode=mode)

    def test_pubmed(self):
        from litsearch import pubmed
        self.assertEqual(pubmed.build_term(self.p("exact")), '"B2 Protein"[Title/Abstract]')
        self.assertEqual(pubmed.build_term(self.p("all")), "(B2[Title/Abstract] AND Protein[Title/Abstract])")
        self.assertEqual(pubmed.build_term(self.p("any")), "(B2[Title/Abstract] OR Protein[Title/Abstract])")

    def test_europepmc_arxiv_epo_uspto(self):
        from litsearch import arxiv, epo, europepmc, patents
        self.assertEqual(europepmc.build_query(self.p("exact")), '(TITLE:"B2 Protein" OR ABSTRACT:"B2 Protein")')
        self.assertIn("(TITLE:B2 OR ABSTRACT:B2) AND (TITLE:Protein OR ABSTRACT:Protein)", europepmc.build_query(self.p("all")))
        self.assertEqual(arxiv.build_query(self.p("exact")), '(ti:"B2 Protein" OR abs:"B2 Protein")')
        self.assertEqual(epo.build_cql(self.p("exact")), '(ti="B2 Protein" or ab="B2 Protein")')
        self.assertIn('ti all "B2 Protein"', epo.build_cql(self.p("all")))
        self.assertIn('ti any "B2 Protein"', epo.build_cql(self.p("any")))
        self.assertEqual(patents.build_q(self.p("exact")), '"B2 Protein"')
        self.assertEqual(patents.build_q(self.p("all")), "(B2 AND Protein)")
        self.assertEqual(patents.build_q(self.p("any")), "(B2 OR Protein)")

    def test_openalex_uses_title_and_abstract_search(self):
        from litsearch import openalex
        captured = {}

        def fake_get(url, params=None, **kw):
            captured["filter"] = params["filter"]
            return types.SimpleNamespace(json=lambda: {"results": [], "meta": {"count": 0}})
        orig, openalex.get = openalex.get, fake_get
        try:
            openalex.search(self.p("exact"), {})
            self.assertEqual(captured["filter"], 'title_and_abstract.search:"B2 Protein"')
            openalex.search(self.p("all"), {})
            self.assertEqual(captured["filter"], "title_and_abstract.search:(B2 AND Protein)")
        finally:
            openalex.get = orig

    def test_one_word_keyword_is_never_quoted(self):
        from litsearch import pubmed
        self.assertEqual(pubmed.build_term(SearchParams(keyword="CRISPR")), "CRISPR[Title/Abstract]")

    def test_stop_words_are_dropped_from_all_and_any(self):
        from litsearch import pubmed
        t = pubmed.build_term(SearchParams(keyword="role of RNA", match_mode="all"))
        self.assertEqual(t, "(role[Title/Abstract] AND RNA[Title/Abstract])")


class FilterAndScoreTests(unittest.TestCase):
    def mk(self, title, abstract="", journal=""):
        a = art(title, "10.1/" + title[:6].replace(" ", ""))
        a.abstract, a.journal = abstract, journal
        return a

    def test_exact_phrase_is_enforced_for_crossref_but_not_trusted_blindly_elsewhere(self):
        from litsearch import records
        p = SearchParams(keyword="B2 protein")
        arts = [self.mk("The B2 protein of nodaviruses"), self.mk("Vitamin B2 and a protein")]
        kept, dropped = records.enforce(arts, p, "crossref")
        self.assertEqual((len(kept), dropped), (1, 1))
        kept, dropped = records.enforce(arts, p, "pubmed")
        self.assertEqual((len(kept), dropped), (2, 0))
        p.match_mode = "all"
        kept, dropped = records.enforce(arts, p, "crossref")
        self.assertEqual(len(kept), 2)

    def test_must_also_contain_and_exclude_with_counts(self):
        from litsearch import records
        p = SearchParams(keyword="B2", must_terms="protein, viral", exclude_terms="riboflavin, HER2")
        arts = [
            self.mk("B2 protein binds dsRNA"),
            self.mk("Riboflavin (vitamin B2) protein intake"),       # exclude by title
            self.mk("HER2 and B2 protein levels"),                    # exclude by title
            self.mk("B2 in milk", abstract="no relevant words"),      # missing must
            self.mk("A viral B2 study", abstract="mentions riboflavin once"),  # abstract-only: kept
        ]
        kept, rules = records.filter_records(arts, p)
        self.assertEqual([a.title for a in kept], ["B2 protein binds dsRNA", "A viral B2 study"])
        self.assertEqual(dict(rules), {"Must also contain: none of the terms found": 1,
                                       "Exclude: riboflavin": 1, "Exclude: HER2": 1})

    def test_terms_match_whole_words_only(self):
        from litsearch.textutil import has_term
        self.assertTrue(has_term("The B2 protein", "B2 protein"))
        self.assertTrue(has_term("The B2-protein", "B2 protein"))
        self.assertFalse(has_term("HER23 levels", "HER2"))
        self.assertFalse(has_term("antibody", "body"))

    def test_score_orders_title_phrase_over_abstract_over_unrelated(self):
        from litsearch import records
        p = SearchParams(keyword="B2 protein", must_terms="viral", exclude_terms="vitamin")
        t = self.mk("The B2 protein of a virus", abstract="A viral protein B2 binds RNA")
        ab = self.mk("Binding of RNA by nodavirus", abstract="We study the B2 protein of a viral family")
        other = self.mk("Cheese ripening", abstract="a vitamin study of B2 proteins")
        s = [records.score(x, p) for x in (t, ab, other)]
        self.assertGreater(s[0], s[1])
        self.assertGreater(s[1], s[2])
        self.assertTrue(all(0 <= v <= 100 for v in s))
        self.assertLess(s[2], 30)                     # abstract-only exclusion hit lowers it

    def test_score_is_neutral_without_a_topic_and_capped(self):
        from litsearch import records
        self.assertEqual(records.score(self.mk("anything"), SearchParams(author="Smith")), 50)
        p = SearchParams(keyword="alpha", must_terms="a, b, c, d", journal="J")
        a = self.mk("alpha", abstract="alpha a b c d", journal="J")
        self.assertLessEqual(records.score(a, p), 100)

    def test_min_score_removes_low_rows_and_logs_it(self):
        orig = dict(search.SOURCES)
        tmp = tempfile.mkdtemp()
        try:
            search.SOURCES.clear()
            search.SOURCES["pubmed"] = ("PubMed", fake([art("Alpha beta protein", "10.1/a"), art("Cheese ripening", "10.1/c")]))
            m = search.run_multi_search(SearchParams(keyword="alpha beta", min_score=30), ["pubmed"],
                                        {"excel_path": os.path.join(tmp, "o.xlsx")})
            self.assertEqual(m.kept, 1)
            self.assertEqual(m.rules, [("Minimum score: below 30", 1)])
        finally:
            search.SOURCES.clear(); search.SOURCES.update(orig)


class ProfileTests(unittest.TestCase):
    def test_examples_are_seeded_once_labelled_and_deletable(self):
        from litsearch import profiles
        cfg = {}
        self.assertTrue(profiles.seed_examples(cfg))
        names = profiles.names(cfg)
        self.assertEqual(names, ["Example: Ambiguous abbreviation", "Example: Patent landscape"])
        self.assertTrue(profiles.delete(cfg, "Example: Patent landscape"))
        self.assertFalse(profiles.seed_examples(cfg))                # stays deleted
        self.assertEqual(profiles.names(cfg), ["Example: Ambiguous abbreviation"])

    def test_save_load_delete_roundtrip(self):
        from litsearch import profiles
        cfg = {}
        profiles.save(cfg, "  My   topic ", {"keyword": "x", "match_mode": "all", "must_terms": "a",
                                              "exclude_terms": "b", "sources": ["pubmed"], "min_score": 5})
        got = profiles.get(cfg, "My topic")
        self.assertEqual((got["keyword"], got["match_mode"], got["sources"]), ("x", "all", ["pubmed"]))
        with self.assertRaises(ValueError):
            profiles.save(cfg, "  ", {})
        self.assertTrue(profiles.delete(cfg, "My topic"))
        self.assertEqual(profiles.get(cfg, "My topic"), {})

    def test_examples_are_generic(self):
        from litsearch import profiles
        text = repr(profiles.EXAMPLES).lower()
        for banned in ("@", "http", "/users/", "api_key"):
            self.assertNotIn(banned, text)


class ConfigTests(unittest.TestCase):
    def test_contact_email_blank_by_default_and_migrates_the_old_setting(self):
        from litsearch import config
        self.assertEqual(config.DEFAULTS["contact_email"], "")
        self.assertEqual(config.DEFAULTS["ncbi_email"], "")
        tmp = Path(tempfile.mkdtemp()) / "c.json"
        tmp.write_text('{"ncbi_email": "someone@example.org"}', encoding="utf-8")
        orig, config.CONFIG_PATH = config.CONFIG_PATH, tmp
        try:
            self.assertEqual(config.load_config()["contact_email"], "someone@example.org")
        finally:
            config.CONFIG_PATH = orig

    def test_polite_pool_email_is_used_only_when_set(self):
        from litsearch import crossref
        seen = []

        def fake_get(url, params=None, **kw):
            seen.append(dict(params))
            return types.SimpleNamespace(json=lambda: {"message": {"items": [], "total-results": 0}})
        orig, crossref.get = crossref.get, fake_get
        try:
            crossref.search(SearchParams(keyword="x"), {})
            crossref.search(SearchParams(keyword="x"), {"contact_email": "me@example.org"})
        finally:
            crossref.get = orig
        self.assertNotIn("mailto", seen[0])
        self.assertEqual(seen[1]["mailto"], "me@example.org")

    def test_no_personal_data_in_the_source(self):
        import re
        root = Path(__file__).resolve().parent.parent
        pat = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
        allowed = {"example.org", "example.com"}
        for p in list(root.glob("*.py")) + list((root / "litsearch").glob("*.py")):
            for m in pat.findall(p.read_text(encoding="utf-8")):
                self.assertIn(m.split("@")[1], allowed, f"{p.name}: {m}")
            self.assertNotIn("/Users/", p.read_text(encoding="utf-8"), p.name)


class MissingKeyAndAiTests(unittest.TestCase):
    def test_patents_and_epo_without_keys_are_skipped_with_a_warning(self):
        tmp = tempfile.mkdtemp()
        says = []
        orig = dict(search.SOURCES)
        try:
            search.SOURCES.clear()
            search.SOURCES["pubmed"] = ("PubMed", fake([art("Alpha paper", "10.1/a")]))
            search.SOURCES["patents"] = ("USPTO Patents", fake([]))
            search.SOURCES["epo"] = ("EPO", fake([]))
            m = search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed", "patents", "epo"],
                                        {"excel_path": os.path.join(tmp, "o.xlsx")}, progress=says.append)
        finally:
            search.SOURCES.clear(); search.SOURCES.update(orig)
        self.assertTrue(m.ok)
        self.assertEqual([r.outcome.skipped for r in m.results], [False, True, True])
        self.assertTrue(any(s.startswith("WARNING") and "USPTO" in s for s in says))
        self.assertTrue(any(s.startswith("WARNING") and "EPO" in s for s in says))

    def test_ai_off_by_default_and_a_missing_ollama_gives_one_line_and_the_search_continues(self):
        from litsearch import ai, config
        self.assertFalse(config.DEFAULTS["ai_enabled"])
        tmp = tempfile.mkdtemp()
        orig = dict(search.SOURCES)
        try:
            search.SOURCES.clear()
            search.SOURCES["pubmed"] = ("PubMed", fake([art("Alpha paper", "10.1/a")]))
            cfg = {"excel_path": os.path.join(tmp, "o.xlsx"), "ai_enabled": True,
                   "ollama_url": "http://localhost:9"}            # nothing listens on port 9
            m = search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], cfg)
        finally:
            search.SOURCES.clear(); search.SOURCES.update(orig)
        self.assertTrue(m.ok)
        self.assertEqual(m.new_rows, 1)
        self.assertEqual(len(m.notes), 1)
        self.assertIn("Ollama is not running", m.notes[0])

    def test_ai_triage_against_a_fake_local_ollama(self):
        """The whole request path with a stand-in server on this computer."""
        import json as _json
        import threading
        from http.server import BaseHTTPRequestHandler, HTTPServer
        from litsearch import ai
        seen = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, obj):
                body = _json.dumps(obj).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self):
                self._send({"models": [{"name": "testmodel:latest"}]})

            def do_POST(self):
                req = _json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                seen.append(req)
                text = req["messages"][1]["content"]
                rating = "Relevant" if "B2 protein" in text.split("Title:")[1] else "Not relevant"
                self._send({"message": {"content": _json.dumps({"rating": rating, "reason": "stub"})}})

        srv = HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            cfg = {"ollama_url": f"http://127.0.0.1:{srv.server_port}", "ollama_model": "testmodel"}
            a, b = art("The B2 protein binds RNA", "10.1/a"), art("Cheese ripening", "10.1/b")
            lines = []
            res = ai.triage([a, b], "Papers about the B2 protein", cfg, lines.append)
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertEqual(res.note, "")
        self.assertEqual((res.model, res.how, res.rated), ("testmodel:latest", "set in settings", 2))
        self.assertEqual((a.ai_rating, b.ai_rating), ("Relevant", "Not relevant"))
        self.assertEqual(a.ai_reason, "stub")
        self.assertEqual(seen[0]["model"], "testmodel:latest")
        self.assertEqual(seen[0]["options"]["temperature"], 0)
        self.assertIn("Papers about the B2 protein", seen[0]["messages"][1]["content"])
        self.assertTrue(any("AI triage" in l for l in lines))

    def _fake_ollama(self, models):
        """A stand-in Ollama on this computer listing ``models``; returns (server, requests seen)."""
        import json as _json
        import threading
        from http.server import BaseHTTPRequestHandler, HTTPServer
        seen = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, obj):
                body = _json.dumps(obj).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self):
                self._send({"models": models})

            def do_POST(self):
                seen.append(_json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                self._send({"message": {"content": '{"rating": "Maybe", "reason": "ok"}'}})

        srv = HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        return srv, seen

    def test_a_blank_model_setting_auto_detects_an_installed_chat_model(self):
        from litsearch import ai, config
        self.assertEqual(config.DEFAULTS["ollama_model"], "")
        srv, seen = self._fake_ollama([
            {"name": "nomic-embed-text:latest", "details": {"family": "nomic-bert"}},
            {"name": "llamacpp:" + "c9" * 32, "details": {"family": "qwen35"}},      # raw alias: skipped
            {"name": "qwen3.5:9b", "details": {"family": "qwen35"}},
            {"name": "older:7b", "details": {"family": "llama"}},
        ])
        try:
            cfg = {"ollama_url": f"http://127.0.0.1:{srv.server_port}"}        # no model named
            a = art("A paper", "10.1/a")
            res = ai.triage([a], "anything", cfg, lambda m: None)
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertEqual((res.model, res.how, res.note), ("qwen3.5:9b", "auto-detected", ""))
        self.assertEqual(seen[0]["model"], "qwen3.5:9b")      # not the embedding model, not a hard-coded default
        self.assertEqual(a.ai_rating, "Maybe")

    def test_a_named_model_that_is_not_installed_falls_back_with_a_note(self):
        from litsearch import ai
        srv, seen = self._fake_ollama([{"name": "other:9b", "details": {"family": "x"}}])
        try:
            cfg = {"ollama_url": f"http://127.0.0.1:{srv.server_port}", "ollama_model": "llama3.1"}
            lines = []
            res = ai.triage([art("A paper", "10.1/a")], "anything", cfg, lines.append)
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertEqual(res.model, "other:9b")
        self.assertTrue(any("'llama3.1' is not installed" in l and "'other:9b' is used instead" in l for l in lines))
        self.assertEqual(seen[0]["model"], "other:9b")

    def test_no_chat_model_installed_gives_a_one_line_note(self):
        from litsearch import ai
        srv, _seen = self._fake_ollama([{"name": "all-minilm", "details": {"family": "bert"}}])
        try:
            res = ai.triage([art("A paper", "10.1/a")], "x", {"ollama_url": f"http://127.0.0.1:{srv.server_port}"}, lambda m: None)
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertEqual(res.model, "")
        self.assertIn("no chat model installed", res.note)

    def test_the_model_used_is_recorded_in_the_search_log(self):
        srv, _seen = self._fake_ollama([{"name": "qwen3.5:9b", "details": {"family": "qwen35"}}])
        tmp = tempfile.mkdtemp()
        orig = dict(search.SOURCES)
        try:
            search.SOURCES.clear()
            search.SOURCES["pubmed"] = ("PubMed", fake([art("Alpha paper", "10.1/a")]))
            cfg = {"excel_path": os.path.join(tmp, "o.xlsx"), "ai_enabled": True,
                   "ollama_url": f"http://127.0.0.1:{srv.server_port}"}
            m = search.run_multi_search(SearchParams(keyword="alpha"), ["pubmed"], cfg)
        finally:
            search.SOURCES.clear(); search.SOURCES.update(orig)
            srv.shutdown()
            srv.server_close()
        self.assertIn("qwen3.5:9b (auto-detected)", m.ai_model)
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(tmp, "o.xlsx"))["Search Log"]
        text = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
        self.assertIn("Local model: qwen3.5:9b (auto-detected)", text)

    def test_only_a_local_ollama_is_ever_used(self):
        from litsearch import ai
        note = ai.check({"ollama_url": "https://api.example.com"})
        self.assertIn("only talks to a local Ollama", note)

    def test_ai_answers_are_parsed_strictly(self):
        from litsearch import ai
        self.assertEqual(ai._parse('{"rating": "Relevant", "reason": "About B2."}'), ("Relevant", "About B2."))
        self.assertEqual(ai._parse('{"rating": "not relevant", "reason": "x"}')[0], "Not relevant")
        self.assertEqual(ai._parse('noise {"rating": "Maybe", "reason": "?"} noise')[0], "Maybe")
        self.assertEqual(ai._parse("garbage")[0], "")


class GuiTests(unittest.TestCase):
    """The window itself (skipped where there is no display)."""

    def setUp(self):
        import litsearch.config as config
        self._orig_cfg = config.CONFIG_PATH
        config.CONFIG_PATH = Path(tempfile.mkdtemp()) / "cfg.json"      # never the real settings
        try:
            from tkinter import Tk
            self.root = Tk()
        except Exception as exc:                                         # no display
            config.CONFIG_PATH = self._orig_cfg
            self.skipTest(f"no display: {exc}")
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
        import app
        self.app_mod = app
        self.app = app.App(self.root)

    def tearDown(self):
        import litsearch.config as config
        try:
            self.root.destroy()
        finally:
            config.CONFIG_PATH = self._orig_cfg

    def test_help_button_opens_the_guide_pdf(self):
        opened = []
        self.app._os_open = lambda p, reveal=False: opened.append(p)
        self.app.open_guide()
        self.assertEqual(len(opened), 1)
        self.assertEqual(opened[0].name, "Litwell_How_To_Guide.pdf")
        self.assertTrue(opened[0].exists())
        self.assertGreater(opened[0].stat().st_size, 50_000)

    def test_about_shows_credit_and_license(self):
        shown = []
        orig = self.app_mod.messagebox.showinfo
        self.app_mod.messagebox.showinfo = lambda title, msg: shown.append(msg)
        try:
            self.app._about()
        finally:
            self.app_mod.messagebox.showinfo = orig
        text = shown[0]
        self.assertIn("Copyright (c) 2026 J. Kent Morgan Consulting, LLC", text)
        self.assertIn("Apache License", text)
        self.assertIn("1.4.3", text)
        self.assertIn("Created by J. Kent Morgan, Ph.D.", text)

    def test_example_profiles_are_in_the_window_and_load(self):
        self.assertEqual(len(self.app.profile_box.cget("values")), 2)
        self.app.profile_name.set("Example: Patent landscape")
        self.app.load_profile()
        self.assertEqual(self.app.match_mode.get(), "all")
        self.assertEqual(self.app._ticked_sources(), ["openalex", "patents", "epo"])


class CrossrefDateTests(unittest.TestCase):
    """Crossref sends [[null]] for records with no date; that crashed the whole search."""

    def test_missing_and_partial_dates(self):
        from litsearch.crossref import _parse_date
        self.assertEqual(_parse_date({"date-parts": [[2024, 5, 17]]}), "2024-05-17")
        self.assertEqual(_parse_date({"date-parts": [[2024, 5]]}), "2024-05")
        self.assertEqual(_parse_date({"date-parts": [[2024]]}), "2024")
        self.assertEqual(_parse_date({"date-parts": [[None]]}), "")
        self.assertEqual(_parse_date({"date-parts": [[2024, None]]}), "2024")
        self.assertEqual(_parse_date({"date-parts": [[]]}), "")
        self.assertEqual(_parse_date(None), "")


if __name__ == "__main__":
    unittest.main()
